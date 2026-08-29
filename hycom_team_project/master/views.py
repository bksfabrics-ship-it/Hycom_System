import logging

from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
import json

from urllib3 import request
from .models import InvoiceSetting, NotificationSetting, Order, OrderItem
from stock.models import Product
from django.shortcuts import render, redirect
from django.forms import formset_factory, modelformset_factory
from .forms import OrderForm, OrderItemForm
from django.contrib import messages
from django.shortcuts import get_object_or_404
from django.http import JsonResponse
from django.db.models import Q, Prefetch, Sum, F, Count
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
import csv
from django.http import HttpResponse
from decimal import Decimal
from returns.models import ReturnItem
from django.forms import modelformset_factory
from django.db import transaction
from datetime import date
from returns.models import Return
from utils.google_sheets import push_order_to_sheet
from utils.email_service import send_order_email
import threading
from openpyxl import Workbook
from django.shortcuts import render
from datetime import datetime
from django.utils import timezone
logger = logging.getLogger(__name__)
from utils.permissions import area_required
from utils.permissions import can_access_area
from .amazon_import import import_amazon_orders
from .portal_imports.amazon_sp_api import import_amazon_orders_via_api
from .portal_imports.flipkart_api import import_flipkart_orders_via_api


def run_async(func, *args):
    thread = threading.Thread(target=func, args=args, daemon=True)
    thread.start()



def update_stock_for_order(order, reverse=False):
    """
    reverse=True → add stock back
    reverse=False → deduct stock
    """

    for item in order.items.all():
        product = item.product

        if reverse:
            product.stock += item.quantity
        else:
            if product.stock < item.quantity:
                raise Exception(f"Not enough stock for {product.name}")
            product.stock -= item.quantity

        product.save()



from django.contrib.auth.decorators import login_required, user_passes_test


def is_staff_user(user):
    return user.is_authenticated and user.is_staff


def staff_or_owner_required(view_func):
    return user_passes_test(is_staff_user, login_url='/accounts/login/')(view_func)


def parse_email_lines(value):
    emails = []
    invalid = []

    for raw_email in (value or "").replace(",", "\n").splitlines():
        email = raw_email.strip().lower()
        if not email:
            continue

        try:
            validate_email(email)
        except ValidationError:
            invalid.append(email)
            continue

        if email not in emails:
            emails.append(email)

    return emails, invalid


def sync_notification_emails(category, emails):
    NotificationSetting.objects.filter(category=category).exclude(email__in=emails).delete()

    for email in emails:
        NotificationSetting.objects.update_or_create(
            category=category,
            email=email,
            defaults={'is_active': True},
        )


def get_notification_emails_text(category):
    return "\n".join(
        NotificationSetting.objects.filter(
            category=category,
            is_active=True
        ).order_by('email').values_list('email', flat=True)
    )


@area_required('settings')
def app_settings(request):
    invoice_settings = InvoiceSetting.load()

    if request.method == 'POST':
        order_update_emails, invalid_order_emails = parse_email_lines(
            request.POST.get('order_update_emails')
        )
        invoice_emails, invalid_invoice_emails = parse_email_lines(
            request.POST.get('invoice_emails')
        )
        invalid_emails = invalid_order_emails + invalid_invoice_emails

        invoice_settings.company_name = request.POST.get('company_name', '').strip() or 'Hycom'
        invoice_settings.company_subtitle = request.POST.get('company_subtitle', '').strip()
        invoice_settings.company_address = request.POST.get('company_address', '').strip()
        invoice_settings.company_phone = request.POST.get('company_phone', '').strip()
        invoice_settings.company_email = request.POST.get('company_email', '').strip()
        invoice_settings.company_gstin = request.POST.get('company_gstin', '').strip()
        invoice_settings.company_pan = request.POST.get('company_pan', '').strip()
        invoice_settings.company_cin = request.POST.get('company_cin', '').strip()
        invoice_settings.company_msme = request.POST.get('company_msme', '').strip()
        invoice_settings.footer_note = request.POST.get('footer_note', '').strip()
        invoice_settings.signature_label = request.POST.get('signature_label', '').strip() or 'Authorised Signature'
        invoice_settings.dc_doc_prefix = request.POST.get('dc_doc_prefix', '').strip() or 'DC'
        invoice_settings.dc_delivery_prefix = request.POST.get('dc_delivery_prefix', '').strip() or 'DN'
        invoice_settings.dc_default_hsn = request.POST.get('dc_default_hsn', '').strip()
        invoice_settings.dc_terms = request.POST.get('dc_terms', '').strip()
        invoice_settings.dc_transport_mode = request.POST.get('dc_transport_mode', '').strip()
        invoice_settings.dc_vehicle_no = request.POST.get('dc_vehicle_no', '').strip()
        invoice_settings.dc_insurance_note = request.POST.get('dc_insurance_note', '').strip()
        invoice_settings.dc_footer_note = request.POST.get('dc_footer_note', '').strip()

        try:
            invoice_settings.full_clean()
        except ValidationError as exc:
            for field_errors in exc.message_dict.values():
                for error in field_errors:
                    messages.error(request, error)
        else:
            if invalid_emails:
                messages.error(
                    request,
                    "Please fix invalid email address(es): " + ", ".join(invalid_emails)
                )
            else:
                invoice_settings.save()
                sync_notification_emails(
                    NotificationSetting.CATEGORY_ORDER_UPDATE,
                    order_update_emails
                )
                sync_notification_emails(
                    NotificationSetting.CATEGORY_INVOICE,
                    invoice_emails
                )
                messages.success(request, "Settings saved successfully.")
                return redirect('/api/settings/')

    context = {
        'invoice_settings': invoice_settings,
        'order_update_emails': get_notification_emails_text(
            NotificationSetting.CATEGORY_ORDER_UPDATE
        ),
        'invoice_emails': get_notification_emails_text(
            NotificationSetting.CATEGORY_INVOICE
        ),
    }

    return render(request, 'settings.html', context)


@area_required('orders')
def get_orders(request):
    if request.method == 'GET':

        orders_data = []

        orders = Order.objects.all().prefetch_related('items')

        for order in orders:
            items = []

            for item in order.items.all():
                items.append({
                    "product": item.product.name,
                    "sku": item.product.sku,
                    "quantity": item.quantity,
                    "price": float(item.price)
                })

            orders_data.append({
                "order_id": order.order_number,
                "customer": order.customer_name,
                "portal": order.portal,
                "status": order.status,
                "total_amount": float(order.total_amount),
                "items": items,
            })
        return JsonResponse(orders_data, safe=False)



# def create_order(request):

#     if request.method == 'POST':
#         order_form = OrderForm(request.POST)
#         formset = OrderItemForm(request.POST)

#         if order_form.is_valid() and formset.is_valid():

#             try:
#                 with transaction.atomic():

#                     order = order_form.save()

#                     items = formset.save(commit=False)

#                     for item in items:
#                         product = item.product
#                         qty = item.quantity

#                         # ❌ VALIDATION: prevent negative stock
#                         if product.stock < qty:
#                             raise Exception(f"Not enough stock for {product.name}")

#                         # ✅ Reduce stock
#                         product.stock -= qty
#                         product.save()

#                         item.order = order
#                         item.save()
#                 messages.success(request, "Order saved successfully")
#                 return redirect('/orders/')

#             except Exception as e:
#                 messages.error(request, "Something went wrong")
#                 return render(request, 'create_order.html', {
#                     'order_form': order_form,
#                     'formset': formset,
#                     'error': str(e)
#                 })

#     else:
#         order_form = OrderForm()
#         formset = OrderItemForm()

#     return render(request, 'create_order.html', {
#         'order_form': order_form,
#         'formset': formset
#     })
        
        

def restore_old_stock(order):
    old_items = OrderItem.objects.filter(order=order)

    for item in old_items:
        product = item.product
        product.stock += item.quantity
        product.save()

    old_items.delete()



@area_required('orders')
def edit_order(request, pk):

    order_model = Order
    order = get_object_or_404(Order, pk=pk)

    ItemFormSet = modelformset_factory(
        OrderItem,
        form=OrderItemForm,
        extra=0,
        can_delete=True
    )

    queryset = OrderItem.objects.filter(order=order)

    if request.method == 'POST':

        order_form = OrderForm(request.POST, instance=order)
        formset = ItemFormSet(request.POST, queryset=queryset)

        if order_form.is_valid() and formset.is_valid():

            try:
                with transaction.atomic():

                    old_status = order_model.objects.filter(id=order.id).values_list('status', flat=True).first()
                    # old_ship_date = order.ship_date

                    # ✅ SAVE ORDER (commit=False)
                    order = order_form.save(commit=False)

                    # ✅ SHIP DATE FIX
                    # if not order.ship_date:
                    #     order.ship_date = old_ship_date

                    # if order.status == 'Shipped' and not old_ship_date:
                    #     order.ship_date = date.today()
                    if order.status == 'Shipped' and not order.ship_date:
                        messages.error(request, "Please select Ship Date when status is Shipped")
                        return redirect(request.path)

                    # if order.status == 'Return In Transit':
                        
                    # ✅ REPLACEMENT LOGIC
                    if order.status == 'Return Arrived':

                        is_replacement = (request.POST.get("is_replacement") or "").lower()

                        if is_replacement == "yes":

                            replacement_id = (request.POST.get("replacement_order_id") or "").strip()

                            if not replacement_id:
                                raise Exception("Replacement Order ID is required")

                            try:
                                replacement_order = Order.objects.get(order_number=replacement_id)
                            except Order.DoesNotExist:
                                raise Exception("Invalid Replacement Order ID")

                            if replacement_order.id == order.id:
                                raise Exception("Order cannot replace itself")

                            order.is_replacement = True
                            order.replacement_for = replacement_order

                        else:
                            order.is_replacement = False
                            order.replacement_for = None

                    else:
                        order.is_replacement = False
                        order.replacement_for = None

                    # ✅ SAVE ORDER ONCE
                    order.save()

                    # ✅ HANDLE ITEMS (NO STOCK LOGIC HERE)
                    items = formset.save(commit=False)

                    for obj in formset.deleted_objects:
                        obj.delete()

                    for item in items:
                        product = item.product
                        price = item.price or product.selling_price

                        item.order = order
                        item.price = price
                        item.save()
                    
                    if old_status != order.status:
                        try:
                            transaction.on_commit(lambda: run_async(send_order_email, order, True))
                        except Exception as e:
                            print("Email failed:", e)
                    
                    try:
                        transaction.on_commit(lambda: run_async(push_order_to_sheet, order))
                    except Exception as e:
                        import traceback
                        print("GOOGLE SYNC ERROR:")
                        traceback.print_exc()
                        
                    already_processed = Return.objects.filter(order=order).exists()
                    if order.status in ['Return Arrived', 'Cancelled']:
                        if already_processed:
                            messages.warning(request, "Return/Cancellation already processed")
                            return redirect('/api/order_list/')
                        return redirect(f'/returns/create/?order_id={order.id}')

                    messages.success(request, "Order updated successfully")
                    user = getattr(request.user, 'username', 'Anonymous') if hasattr(request, 'user') and request.user.is_authenticated else 'Anonymous'
                    logger.info(f"Order updated: {order.order_number} (status: {order.status}) by user {user}")
                    return redirect('/api/order_list/')

            except Exception as e:
                messages.error(request, str(e))

        else:
            print("ORDER ERR:", order_form.errors)
            print("FORMSET ERR:", formset.errors)
            messages.error(request, "Please fix form errors")

    else:
        order_form = OrderForm(instance=order)
        formset = ItemFormSet(queryset=queryset)

    return render(request, 'edit_order.html', {
        'order_form': order_form,
        'formset': formset,
        'order': order
    })

def get_product_by_sku(request):
    sku = request.GET.get('sku', '').strip()

    try:
        product = Product.objects.get(sku__iexact=sku)  # ✅ case-insensitive

        return JsonResponse({
            'id': product.id,
            'name': product.name,
            'price': float(product.selling_price)
        })
        print("SKU RECEIVED:", sku)
    except Product.DoesNotExist:
        return JsonResponse({'error': 'Product not found'})



@area_required('products')
def create_order_ui(request):

    ItemFormSet = formset_factory(OrderItemForm, extra=1)

    if request.method == 'POST':
        order_form = OrderForm(request.POST)
        formset = ItemFormSet(request.POST)

        if order_form.is_valid() and formset.is_valid():

            try:
                with transaction.atomic():

                    # ✅ STEP 1: Create order (NOT saved fully yet)
                    order = order_form.save(commit=False)
                    order.amount = 0
                    order.gst = 0
                    order.cgst = 0
                    order.sgst = 0
                    order.igst = 0
                    order.total_amount = 0
                    order.save()   # MUST save before using FK

                    valid_items = 0

                    # ✅ STEP 2: Save items with validation
                    for form in formset:
                        data = form.cleaned_data

                        if not data or not data.get('product') or not data.get('quantity'):
                            continue

                        product = data['product']
                        qty = data['quantity']
                        price = data.get('price') or product.selling_price

                        # 🔴 STOCK VALIDATION
                        if not order.is_stock_updated:
                            if product.stock < qty:
                                raise Exception(f"Not enough stock for {product.name}")

                            # ✅ Reduce stock
                            product.stock -= qty
                            product.save()

                        item = form.save(commit=False)
                        item.order = order
                        item.price = price
                        item.save()

                        valid_items += 1

                    # ❌ No items
                    if valid_items == 0:
                        raise Exception("At least one valid item is required")

                    # ✅ STEP 3: Calculate totals
                    amount = Decimal('0')
                    gst_total = Decimal('0')

                    for it in order.items.all():
                        total = it.quantity * it.price  # already Decimal

                        base = total / Decimal('1.05')
                        gst = total - base

                        amount += base
                        gst_total += gst

                    order.amount = round(amount, 2)
                    order.gst = round(gst_total, 2)

                    # ✅ STEP 4: GST Split
                    print("STATE CODE:,,,,,,,", order.state.code if order.state else "NO STATE")
                    state_code = (order.state.code if order.state and order.state.code else "")
                    print("STATE CODE AFTER STRIP:,,,,,,,", state_code)
                    if state_code == '33':
                        print("INTRA-STATE ORDER - APPLYING CGST/SGST SPLIT")
                        order.cgst = round(order.gst / 2, 2)
                        order.sgst = round(order.gst / 2, 2)
                        order.igst = 0
                    else:
                        order.igst = round(order.gst, 2)
                        order.cgst = 0
                        order.sgst = 0

                    # ✅ STEP 5: Final total
                    order.total_amount = round(order.amount + order.gst, 2)

                    order.save()
                    
                    
                    try:
                        transaction.on_commit(lambda: run_async(send_order_email, order, True))
                    except Exception as e:
                        print("Email failed:", e)
                    
                    try:
                        transaction.on_commit(lambda: run_async(push_order_to_sheet, order))
                    except Exception as e:
                        import traceback
                        print("GOOGLE SYNC ERROR:")
                        traceback.print_exc()
                        
                messages.success(request, "Order saved successfully")
                user = getattr(request.user, 'username', 'Anonymous') if hasattr(request, 'user') and request.user.is_authenticated else 'Anonymous'
                logger.info(f"Order created: {order.order_number} by user {user}")
                return redirect('/api/order_list/')

            except Exception as e:
                messages.error(request, f"Error: {str(e)}")

        else:
            messages.error(request, "Please correct the errors below.")

            print("ORDER FORM ERRORS:", order_form.errors)
            print("FORMSET ERRORS:", formset.errors)

            return render(request, 'create_order.html', {
                'order_form': order_form,
                'formset': formset
            })

    else:
        order_form = OrderForm()
        formset = ItemFormSet()

        return render(request, 'create_order.html', {
            'order_form': order_form,
            'formset': formset
        })



    total_products = Product.objects.count()
    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')
    if from_date and to_date:
        total_orders = Order.objects.filter(invoice_date__range=[from_date, to_date]).count()
    else:
        total_orders = Order.objects.count()
    total_stock = Product.objects.aggregate(total=Sum('stock'))['total'] or 0

    # ✅ MASTER LISTS
    ALL_STYLES = ['Core', 'Flexi', 'Ethos']
    ALL_COLORS = ['Wine Red', 'Hunter Green', 'Ceil Blue', 'Navy Blue']
    ALL_GENDERS = ['Male', 'Female']
    ALL_SIZES = ['XS', 'S', 'M', 'L', 'XL', '2XL']

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    items = OrderItem.objects.select_related('product', 'order')

    # ✅ DATE FILTER
    if from_date and to_date:
        items = items.filter(order__invoice_date__range=[from_date, to_date])

    # ✅ TOP PRODUCTS
    top_products = (
        items.values('product__name', 'product__sku')
        .annotate(total_qty=Sum('quantity'))
        .order_by('-total_qty')[:10]
    )

    # RAW QUERYSETS
    style_qs = (
        items.values('product__style')
        .exclude(product__style__isnull=True)
        .exclude(product__style__exact='')
        .annotate(total=Sum('quantity'))
    )

    color_qs = (
        items.values('product__color')
        .exclude(product__color__isnull=True)
        .exclude(product__color__exact='')
        .annotate(total=Sum('quantity'))
    )

    gender_qs = (
        items.values('product__gender')
        .exclude(product__gender__isnull=True)
        .exclude(product__gender__exact='')
        .annotate(total=Sum('quantity'))
    )

    size_qs = (
        items.values('product__size')
        .exclude(product__size__isnull=True)
        .exclude(product__size__exact='')
        .annotate(total=Sum('quantity'))
    )

    # FILL MISSING
    def fill_missing(data, key, all_values):
        data_dict = {item[key]: item['total'] for item in data}
        final = []
        for value in all_values:
            final.append({key: value, 'total': data_dict.get(value, 0)})
        return final

    style_data = fill_missing(list(style_qs), 'product__style', ALL_STYLES)
    color_data = fill_missing(list(color_qs), 'product__color', ALL_COLORS)
    gender_data = fill_missing(list(gender_qs), 'product__gender', ALL_GENDERS)
    size_data = fill_missing(list(size_qs), 'product__size', ALL_SIZES)

    context = {
        'total_products': total_products,
        'total_orders': total_orders,
        'total_stock': total_stock,
        'top_products': top_products,
        'style_data': style_data,
        'color_data': color_data,
        'gender_data': gender_data,
        'size_data': size_data,
    }

    return render(request, 'dashboard.html', context)

    
    


@area_required('orders')
def order_list(request):
    search  = request.GET.get('q')
    status = request.GET.get('status')
    portal = request.GET.get('portal')
    
    if not can_access_area(request.user, 'orders'):
        return render(request, '403.html')

    orders = Order.objects.all().order_by('-id').prefetch_related(Prefetch('items', queryset=OrderItem.objects.select_related('product')))
    
    if status:
        orders = orders.filter(status=status)

    if portal:
        orders = orders.filter(portal=portal)

    if search:
        orders = orders.filter(order_number__icontains=search)

    returns_data = ReturnItem.objects.values('return_obj__order') \
        .annotate(
            return_count=Count('id'),
            return_qty=Sum('quantity')
        )

    returns_map = {
        r['return_obj__order']: r for r in returns_data
    }

    for order in orders:
        data = returns_map.get(order.id, {})
        order.return_count = data.get('return_count', 0)
        order.return_qty = data.get('return_qty', 0)

    return render(request, 'order_list.html', {
        'orders': orders,
    })


@area_required('orders')
def import_orders(request):
    import_result = None
    sp_import_result = None
    flipkart_import_result = None

    if request.method == 'POST':
        import_source = request.POST.get('import_source', 'excel')

        if import_source == 'sp_api':
            from_date_raw = request.POST.get('from_date')
            to_date_raw = request.POST.get('to_date')

            try:
                f_date = datetime.strptime(from_date_raw, "%Y-%m-%d").date()
                t_date = datetime.strptime(to_date_raw, "%Y-%m-%d").date()
                if f_date > t_date:
                    raise ValueError("from_date is after to_date")
            except (ValueError, TypeError):
                messages.error(request, "Invalid date range. Use YYYY-MM-DD.")
            else:
                try:
                    sp_import_result = import_amazon_orders_via_api(f_date, t_date)
                except Exception as exc:
                    messages.error(request, f"SP-API import failed: {exc}")
                else:
                    if sp_import_result.created_count or sp_import_result.updated_count:
                        messages.success(
                            request,
                            f"Amazon SP-API sync complete: "
                            f"{sp_import_result.created_count} created, "
                            f"{sp_import_result.updated_count} updated."
                        )
                    if sp_import_result.skipped_count:
                        messages.warning(
                            request,
                            f"Skipped {sp_import_result.skipped_count} SP-API order(s). Review details below."
                        )
                    if sp_import_result.failed_count:
                        messages.error(
                            request,
                            f"{sp_import_result.failed_count} SP-API order(s) failed. Review details below."
                        )
                    if sp_import_result.errors and not (
                        sp_import_result.created_count or sp_import_result.updated_count
                    ):
                        messages.error(request, "No SP-API orders were imported. Review details below.")
        elif import_source == 'flipkart_api':
            from_date_raw = request.POST.get('from_date')
            to_date_raw = request.POST.get('to_date')

            try:
                f_date = datetime.strptime(from_date_raw, "%Y-%m-%d").date()
                t_date = datetime.strptime(to_date_raw, "%Y-%m-%d").date()
                if f_date > t_date:
                    raise ValueError("from_date is after to_date")
            except (ValueError, TypeError):
                messages.error(request, "Invalid date range. Use YYYY-MM-DD.")
            else:
                try:
                    flipkart_import_result = import_flipkart_orders_via_api(f_date, t_date)
                except Exception as exc:
                    messages.error(request, f"Flipkart import failed: {exc}")
                else:
                    if flipkart_import_result.created_count or flipkart_import_result.updated_count:
                        messages.success(
                            request,
                            f"Flipkart sync complete: "
                            f"{flipkart_import_result.created_count} created, "
                            f"{flipkart_import_result.updated_count} updated."
                        )
                    if flipkart_import_result.skipped_count:
                        messages.warning(
                            request,
                            f"Skipped {flipkart_import_result.skipped_count} Flipkart order(s). Review details below."
                        )
                    if flipkart_import_result.failed_count:
                        messages.error(
                            request,
                            f"{flipkart_import_result.failed_count} Flipkart order(s) failed. Review details below."
                        )
                    if flipkart_import_result.errors and not (
                        flipkart_import_result.created_count or flipkart_import_result.updated_count
                    ):
                        messages.error(request, "No Flipkart orders were imported. Review details below.")
        else:
            report_file = request.FILES.get('order_report')

            if not report_file:
                messages.error(request, "Please choose an Amazon order report Excel file.")
            elif not report_file.name.lower().endswith((".xlsx", ".xlsm")):
                messages.error(request, "Please upload an Excel file in .xlsx or .xlsm format.")
            else:
                try:
                    import_result = import_amazon_orders(report_file)
                except Exception as exc:
                    messages.error(request, f"Import failed: {exc}")
                else:
                    if import_result.created_count:
                        messages.success(
                            request,
                            f"Imported {import_result.created_count} order(s) with "
                            f"{import_result.item_count} item(s)."
                        )
                    if import_result.skipped_count:
                        messages.warning(
                            request,
                            f"Skipped {import_result.skipped_count} order(s). Review details below."
                        )
                    if import_result.errors and not import_result.created_count:
                        messages.error(request, "No orders were imported. Review details below.")

    return render(request, 'import_orders.html', {
        'import_result': import_result,
        'sp_import_result': sp_import_result,
        'flipkart_import_result': flipkart_import_result,
    })


@area_required('orders')
def order_invoice(request, pk):
    order = get_object_or_404(
        Order.objects.select_related('state').prefetch_related(
            Prefetch('items', queryset=OrderItem.objects.select_related('product'))
        ),
        pk=pk
    )

    invoice_items = []
    for index, item in enumerate(order.items.all(), start=1):
        unit_price = item.price or item.product.selling_price or Decimal('0')
        quantity = Decimal(item.quantity or 0)
        invoice_items.append({
            'sl_no': index,
            'product': item.product,
            'quantity': item.quantity,
            'unit_price': unit_price,
            'line_total': unit_price * quantity,
        })

    return render(request, 'invoice.html', {
        'order': order,
        'invoice_items': invoice_items,
        'invoice_settings': InvoiceSetting.load(),
    })


@area_required('orders')
def order_delivery_challan(request, pk):
    order = get_object_or_404(
        Order.objects.select_related('state').prefetch_related(
            Prefetch('items', queryset=OrderItem.objects.select_related('product'))
        ),
        pk=pk
    )

    if not order.dc_date:
        order.dc_date = timezone.localdate()
        order.save(update_fields=['dc_date'])

    dc_items = []
    for index, item in enumerate(order.items.all(), start=1):
        dc_items.append({
            'sl_no': index,
            'product': item.product,
            'quantity': item.quantity,
            'qty_kgs': item.quantity,
            'uom': 'SET',
            'batch': item.product.material_code or item.product.sku,
        })

    settings = InvoiceSetting.load()
    doc_prefix = settings.dc_doc_prefix or 'DC'
    delivery_prefix = settings.dc_delivery_prefix or 'DN'

    return render(request, 'delivery_challan.html', {
        'order': order,
        'dc_items': dc_items,
        'dc_settings': settings,
        'dc_doc_no': f"{doc_prefix}{order.id:06d}",
        'dc_delivery_no': f"{delivery_prefix}{order.id:06d}",
    })
    
    

def delete_order(request, pk):

    order = get_object_or_404(Order, id=pk)

    # 🔁 restore stock
    items = OrderItem.objects.filter(order=order)

    already_restored = ReturnItem.objects.filter(return_obj__order=order).exists()

    if not already_restored:
        for item in items:
            product = item.product
            product.stock += item.quantity
            product.save()

    order.delete()

    messages.success(request, "Order deleted successfully")
    user = getattr(request.user, 'username', 'Anonymous') if hasattr(request, 'user') and request.user.is_authenticated else 'Anonymous'
    logger.info(f"Order deleted: {order.order_number} by user {user}")
    return redirect('/api/order_list/')



def export_orders(request):

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="orders.csv"'

    writer = csv.writer(response)

    # HEADER
    writer.writerow([
        'Order No', 'Customer', 'Invoice No', 'Invoice Date', 'Ship Date',
        'Portal', 'Status', 'State', 'State Code',
        'Qty', 'Amount', 'Fulfilment', 'B2B', 'GST',
        'SKU', 'Material Code'
    ])

    orders = Order.objects.all().prefetch_related('items')

    for o in orders:
        for item in o.items.all():
            writer.writerow([
                o.order_number,
                o.customer_name,
                o.invoice_number,
                o.invoice_date,
                o.ship_date,
                o.portal,
                o.status,
                o.state,
                o.state_code,
                item.quantity,
                item.price,
                o.fulfilment,
                o.is_b2b,
                o.gst_number,
                item.product.sku,
                item.product.material_code
            ])

    return response


def fill_missing(data, key, all_values):
    """
    Ensures all categories exist, fills missing with 0
    """
    data_dict = {item[key]: item['total'] for item in data}

    final = []
    for value in all_values:
        final.append({
            key: value,
            'total': data_dict.get(value, 0)
        })
    return final


@area_required('dashboard')
def dashboard(request):
    from stock.models import Product
    
    if not can_access_area(request.user, 'dashboard'):
        return render(request, '403.html')
    # KPIs
    total_products = Product.objects.count()
    total_orders = Order.objects.count()
    total_stock = Product.objects.aggregate(total=Sum('stock'))['total'] or 0

    # ✅ MASTER LISTS
    ALL_STYLES = ['Core', 'Flexi', 'Ethos']
    ALL_COLORS = ['Wine Red', 'Hunter Green', 'Ceil Blue', 'Navy Blue']
    ALL_GENDERS = ['Men', 'Women']
    ALL_SIZES = ['XS', 'S', 'M', 'L', 'XL', '2XL']

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    items = OrderItem.objects.select_related('product', 'order')

    # ✅ DATE FILTER
    if from_date and to_date:
        items = items.filter(order__invoice_date__range=[from_date, to_date])

    # ✅ TOP PRODUCTS
    top_products = (
        items.values('product__name', 'product__sku')
        .annotate(total_qty=Sum('quantity'))
        .order_by('-total_qty')[:10]
    )

    # RAW QUERYSETS
    style_qs = (
        items.values('product__style')
        .exclude(product__style__isnull=True)
        .exclude(product__style__exact='')
        .annotate(total=Sum('quantity'))
    )

    color_qs = (
        items.values('product__color')
        .exclude(product__color__isnull=True)
        .exclude(product__color__exact='')
        .annotate(total=Sum('quantity'))
    )

    gender_qs = (
        items.values('product__gender')
        .exclude(product__gender__isnull=True)
        .exclude(product__gender__exact='')
        .annotate(total=Sum('quantity'))
    )

    size_qs = (
        items.values('product__size')
        .exclude(product__size__isnull=True)
        .exclude(product__size__exact='')
        .annotate(total=Sum('quantity'))
    )

    # FILL MISSING
    style_data = fill_missing(list(style_qs), 'product__style', ALL_STYLES)
    color_data = fill_missing(list(color_qs), 'product__color', ALL_COLORS)
    gender_data = fill_missing(list(gender_qs), 'product__gender', ALL_GENDERS)
    size_data = fill_missing(list(size_qs), 'product__size', ALL_SIZES)

    context = {
        'total_products': total_products,
        'total_orders': total_orders,
        'total_stock': total_stock,
        'top_products': top_products,
        'style_data': style_data,
        'color_data': color_data,
        'gender_data': gender_data,
        'size_data': size_data,
    }
    

    return render(request, 'dashboard.html', context)


@area_required('dashboard')
def export_dashboard_excel(request):
    from openpyxl.styles import Font

    from_date = request.GET.get('from_date')
    to_date = request.GET.get('to_date')

    total_products = Product.objects.count()
    if from_date and to_date:
        total_orders = Order.objects.filter(invoice_date__range=[from_date, to_date]).count()
    else:
        total_orders = Order.objects.count()
    total_stock = Product.objects.aggregate(total=Sum('stock'))['total'] or 0

    all_styles = ['Core', 'Flexi', 'Ethos']
    all_colors = ['Wine Red', 'Hunter Green', 'Ceil Blue', 'Navy Blue']
    all_genders = ['Male', 'Female', 'Men', 'Women']
    all_sizes = ['XS', 'S', 'M', 'L', 'XL', '2XL']

    items = OrderItem.objects.select_related('product', 'order')
    if from_date and to_date:
        items = items.filter(order__invoice_date__range=[from_date, to_date])

    top_products = (
        items.values('product__name', 'product__sku')
        .annotate(total_qty=Sum('quantity'))
        .order_by('-total_qty')[:10]
    )
    style_data = fill_missing(list(
        items.values('product__style')
        .exclude(product__style__isnull=True)
        .exclude(product__style__exact='')
        .annotate(total=Sum('quantity'))
    ), 'product__style', all_styles)
    color_data = fill_missing(list(
        items.values('product__color')
        .exclude(product__color__isnull=True)
        .exclude(product__color__exact='')
        .annotate(total=Sum('quantity'))
    ), 'product__color', all_colors)
    gender_data = fill_missing(list(
        items.values('product__gender')
        .exclude(product__gender__isnull=True)
        .exclude(product__gender__exact='')
        .annotate(total=Sum('quantity'))
    ), 'product__gender', all_genders)
    size_data = fill_missing(list(
        items.values('product__size')
        .exclude(product__size__isnull=True)
        .exclude(product__size__exact='')
        .annotate(total=Sum('quantity'))
    ), 'product__size', all_sizes)

    wb = Workbook()
    wb.remove(wb.active)

    ws_summary = wb.create_sheet('Summary')
    ws_summary.append(['Metric', 'Value'])
    ws_summary.append(['Total Products', total_products])
    ws_summary.append(['Total Orders', total_orders])
    ws_summary.append(['Total Stock', total_stock])
    ws_summary.append([''])
    ws_summary.append(['Filter Period', f"{from_date or 'All'} to {to_date or 'All'}"])

    ws_top = wb.create_sheet('Top Products')
    ws_top.append(['Product', 'SKU', 'Total Qty'])
    for product in top_products:
        ws_top.append([
            product['product__name'],
            product['product__sku'],
            product['total_qty'],
        ])

    for sheet_name, label, key, rows in [
        ('Style Wise', 'Style', 'product__style', style_data),
        ('Color Wise', 'Color', 'product__color', color_data),
        ('Gender Wise', 'Gender', 'product__gender', gender_data),
        ('Size Wise', 'Size', 'product__size', size_data),
    ]:
        worksheet = wb.create_sheet(sheet_name)
        worksheet.append([label, 'Quantity'])
        total = 0
        for row in rows:
            qty = row['total']
            worksheet.append([row[key], qty])
            total += qty
        worksheet.append(['TOTAL', total])

    bold_font = Font(bold=True)
    for worksheet in wb.worksheets:
        for cell in worksheet[1]:
            cell.font = bold_font
        last_row = worksheet.max_row
        if last_row > 1 and worksheet.cell(row=last_row, column=1).value == 'TOTAL':
            for cell in worksheet[last_row]:
                cell.font = bold_font
        for column in worksheet.columns:
            column_letter = column[0].column_letter
            max_length = max(len(str(cell.value or '')) for cell in column)
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename=dashboard.xlsx'
    wb.save(response)
    return response
