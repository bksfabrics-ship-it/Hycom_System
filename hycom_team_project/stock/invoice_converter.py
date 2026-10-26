from io import BytesIO
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


INPUT_COLUMNS = (
    "DC Number",
    "DC Date",
    "GST (if B2B Customer)",
    "Portal",
    "SAP SKU",
    "Qty",
    "Taxable Amount",
    "Tax",
    "Shipping Charges",
    "Total Amount",
)

COLUMN_ALIASES = {
    "shipping": "shipping charges",
}


def clean_value(value):
    if pd.isna(value):
        return ""
    return str(value).strip().replace("\xa0", " ")


def clean_number(value):
    cleaned = clean_value(value)
    if cleaned == "":
        return None

    try:
        return float(value)
    except (ValueError, TypeError):
        return cleaned


def _normalized_column_name(column):
    return " ".join(str(column).strip().casefold().split())


def _prepare_input_columns(df):
    columns_by_name = {}
    for column in df.columns:
        normalized = _normalized_column_name(column)
        canonical = COLUMN_ALIASES.get(normalized, normalized)
        columns_by_name.setdefault(canonical, column)

    for column in INPUT_COLUMNS:
        source_column = columns_by_name.get(_normalized_column_name(column))
        if source_column is None:
            df[column] = ""
        elif source_column != column:
            df[column] = df[source_column]

    return df


def _worksheet_name(dc_number, existing_names):
    base_name = re.sub(r"[\[\]:*?/\\]", "_", f"DC_{dc_number}").strip("'")
    base_name = base_name[:31] or "DC"
    name = base_name
    suffix_number = 2

    while name.casefold() in existing_names:
        suffix = f"_{suffix_number}"
        name = f"{base_name[:31 - len(suffix)]}{suffix}"
        suffix_number += 1

    existing_names.add(name.casefold())
    return name


def convert_dc_excel(input_file):
    """Convert an uploaded DC input workbook into the SAP invoice workbook."""
    try:
        df = pd.read_excel(input_file)
    except Exception as exc:
        raise ValueError("The uploaded file could not be read as an Excel workbook.") from exc

    df.columns = [str(column).strip() for column in df.columns]
    df = df.dropna(how="all").copy()
    df = _prepare_input_columns(df)
    df["DC Number"] = df["DC Number"].apply(clean_value)
    df["DC Number"] = df["DC Number"].replace("", pd.NA).ffill().fillna("")

    workbook = Workbook()
    workbook.remove(workbook.active)

    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    header_fill = PatternFill(fill_type="solid", fgColor="D9EAD3")
    header_font = Font(name="Calibri", size=11, bold=False)
    normal_font = Font(name="Calibri", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center")

    existing_names = set()
    groups = (
        df.groupby("DC Number", sort=False, dropna=False)
        if not df.empty
        else [("", df)]
    )
    for dc_number, group in groups:
        ws = workbook.create_sheet(_worksheet_name(dc_number, existing_names))

        for column, width in {
            "A": 48,
            "B": 20,
            "C": 12,
            "D": 14,
            "E": 12,
            "F": 15,
        }.items():
            ws.column_dimensions[column].width = width

        ws["C1"] = "Remarks"
        ws["E1"] = "Shipping Charges"
        ws["F1"] = "Total Value"

        for cell in ws[1]:
            cell.fill = header_fill
            cell.border = thin_border
            cell.font = header_font
            cell.alignment = center_alignment
        for cell in ("C1", "E1", "F1"):
            ws[cell].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 25

        first_row = (
            group.iloc[0]
            if not group.empty
            else pd.Series("", index=INPUT_COLUMNS)
        )
        gst_value = clean_value(first_row["GST (if B2B Customer)"]).lower()
        ws["A3"] = "Date of Invoice"
        ws["B3"] = clean_value(first_row["DC Date"])
        ws["A4"] = "Invoice Type"
        ws["B4"] = "B TO B" if gst_value == "yes" else (
            "B TO C" if gst_value == "no" else ""
        )
        ws["A5"] = "Buyer"
        ws["B5"] = clean_value(first_row["Portal"])
        ws["A7"] = "DC no"
        ws["B7"] = dc_number

        for row_number in (3, 4, 5, 7):
            for column_number in range(1, 7):
                cell = ws.cell(row=row_number, column=column_number)
                cell.border = thin_border
                cell.font = normal_font
                cell.alignment = center_alignment

        current_row = 9
        for item_number, (_, row) in enumerate(group.iterrows(), start=1):
            quantity = clean_number(row["Qty"])
            taxable_amount = clean_number(row["Taxable Amount"])
            tax = clean_number(row["Tax"])
            shipping = clean_number(row["Shipping Charges"])
            total_amount = clean_number(row["Total Amount"])
            material_code = clean_value(row["SAP SKU"])

            ws.cell(
                row=current_row,
                column=1,
                value=f"Dispatched Quantity  #{item_number}",
            )
            ws.cell(row=current_row, column=2, value=quantity)
            for column_number in range(1, 7):
                cell = ws.cell(row=current_row, column=column_number)
                cell.border = thin_border
                cell.alignment = center_alignment
            current_row += 1

            ws.cell(
                row=current_row,
                column=1,
                value=f"Dispatched goods against Invoice value # {item_number}",
            )
            ws.cell(row=current_row, column=2, value=taxable_amount)
            ws.cell(row=current_row, column=3, value=5)
            ws.cell(row=current_row, column=4, value=tax)
            ws.cell(row=current_row, column=5, value=shipping)
            ws.cell(row=current_row, column=6, value=total_amount)
            for column_number in range(1, 7):
                cell = ws.cell(row=current_row, column=column_number)
                cell.border = thin_border
                cell.alignment = center_alignment

            for column_number in (2, 4, 5, 6):
                ws.cell(row=current_row, column=column_number).number_format = (
                    "#,##0.00"
                )
            current_row += 1

            ws.cell(
                row=current_row,
                column=1,
                value=f"Material Code # {item_number}",
            )
            ws.cell(row=current_row, column=2, value=material_code)
            for column_number in range(1, 7):
                cell = ws.cell(row=current_row, column=column_number)
                cell.border = thin_border
                cell.alignment = center_alignment
            current_row += 2

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.font = normal_font
                    cell.alignment = center_alignment

        ws.freeze_panes = "A3"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
