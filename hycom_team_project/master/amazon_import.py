from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.core.exceptions import ValidationError
from django.utils.dateparse import parse_datetime
from openpyxl import load_workbook

from stock.models import Product

from .models import Order, OrderItem, State


REQUIRED_HEADERS = {
    "amazon-order-id",
    "purchase-date",
    "order-status",
    "fulfillment-channel",
    "sku",
    "quantity",
    "item-price",
    "ship-state",
}


@dataclass
class ImportRow:
    row_number: int
    data: dict


@dataclass
class ImportResult:
    created_count: int = 0
    skipped_count: int = 0
    item_count: int = 0
    errors: list[str] = field(default_factory=list)
    skipped_orders: list[str] = field(default_factory=list)
    created_orders: list[str] = field(default_factory=list)


def normalize_header(value):
    return str(value or "").strip().lower()


def as_decimal(value):
    if value in (None, ""):
        return Decimal("0")
    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return Decimal("0")


def as_int(value):
    if value in (None, ""):
        return 0
    try:
        return int(Decimal(str(value)))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def as_bool(value):
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "yes", "y"}


def parse_report_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return date.today()

    parsed = parse_datetime(str(value))
    if parsed:
        return parsed.date()

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value), fmt).date()
        except ValueError:
            continue

    return date.today()


def normalize_status(value):
    raw_status = str(value or "").strip().lower()
    if "cancel" in raw_status:
        return "Cancelled"
    if "ship" in raw_status and "unshipped" not in raw_status:
        return "Shipped"
    return "Pending"


def normalize_fulfilment(value):
    channel = str(value or "").strip().lower()
    if channel == "amazon":
        return "amazon"
    if channel == "flipkart":
        return "flipkart"
    return "self"


def find_state(value):
    state_name = str(value or "").strip()
    if not state_name:
        return None

    return State.objects.filter(name__iexact=state_name).first()


def find_product(sku):
    sku = str(sku or "").strip()
    if not sku:
        return None

    product = Product.objects.filter(sku__iexact=sku).first()
    if product:
        return product

    if sku.upper().endswith("_FBA"):
        return Product.objects.filter(sku__iexact=sku[:-4]).first()

    return None


def row_total(row):
    data = row.data
    return (
        as_decimal(data.get("item-price"))
        + as_decimal(data.get("shipping-price"))
        + as_decimal(data.get("gift-wrap-price"))
        - as_decimal(data.get("item-promotion-discount"))
        - as_decimal(data.get("ship-promotion-discount"))
    )


def row_tax(row):
    return (
        as_decimal(row.data.get("item-tax"))
        + as_decimal(row.data.get("shipping-tax"))
        + as_decimal(row.data.get("gift-wrap-tax"))
    )


def load_amazon_rows(uploaded_file):
    workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)

    try:
        headers = [normalize_header(value) for value in next(rows)]
    except StopIteration:
        raise ValueError("The uploaded Excel file is empty.")

    missing_headers = sorted(REQUIRED_HEADERS - set(headers))
    if missing_headers:
        raise ValueError(
            "Missing required column(s): " + ", ".join(missing_headers)
        )

    parsed_rows = []
    for row_number, values in enumerate(rows, start=2):
        data = dict(zip(headers, values))
        if not any(value not in (None, "") for value in data.values()):
            continue
        parsed_rows.append(ImportRow(row_number=row_number, data=data))

    return parsed_rows


def import_amazon_orders(uploaded_file):
    result = ImportResult()
    rows = load_amazon_rows(uploaded_file)

    grouped_rows = defaultdict(list)
    for row in rows:
        order_id = str(row.data.get("amazon-order-id") or "").strip()
        if not order_id:
            result.errors.append(f"Row {row.row_number}: amazon-order-id is missing.")
            continue
        grouped_rows[order_id].append(row)

    for order_number, order_rows in grouped_rows.items():
        if Order.objects.filter(order_number=order_number).exists():
            result.skipped_count += 1
            result.skipped_orders.append(f"{order_number}: already exists")
            continue

        first_row = order_rows[0].data
        invoice_number = str(first_row.get("merchant-order-id") or order_number)
        if Order.objects.filter(invoice_number=invoice_number).exists():
            result.skipped_count += 1
            result.skipped_orders.append(f"{order_number}: invoice number already exists")
            continue

        positive_item_rows = [
            row for row in order_rows if as_int(row.data.get("quantity")) > 0
        ]
        products_by_row = []
        has_product_error = False

        for row in positive_item_rows:
            sku = row.data.get("sku")
            product = find_product(sku)
            if not product:
                result.errors.append(
                    f"Order {order_number}, row {row.row_number}: SKU '{sku}' was not found."
                )
                has_product_error = True
                continue
            products_by_row.append((row, product))

        if has_product_error:
            result.skipped_count += 1
            result.skipped_orders.append(f"{order_number}: missing product SKU")
            continue

        state = find_state(first_row.get("ship-state"))
        total = sum((row_total(row) for row in order_rows), Decimal("0"))
        gst = sum((row_tax(row) for row in order_rows), Decimal("0"))

        if gst == 0 and total > 0:
            base_amount = total / Decimal("1.05")
            gst = total - base_amount
        else:
            base_amount = total - gst

        try:
            with transaction.atomic():
                order = Order.objects.create(
                    portal="amazon",
                    order_number=order_number,
                    customer_name="Amazon Customer",
                    state=state,
                    state_code=state.code if state else "",
                    invoice_number=invoice_number,
                    invoice_date=parse_report_date(first_row.get("purchase-date")),
                    ship_date=parse_report_date(first_row.get("last-updated-date"))
                    if normalize_status(first_row.get("order-status")) == "Shipped"
                    else None,
                    fulfilment=normalize_fulfilment(first_row.get("fulfillment-channel")),
                    is_b2b=as_bool(first_row.get("is-business-order")),
                    status=normalize_status(first_row.get("order-status")),
                    amount=round(base_amount, 2),
                    gst=round(gst, 2),
                    cgst=round(gst / 2, 2)
                    if state and state.code == "33"
                    else Decimal("0"),
                    sgst=round(gst / 2, 2)
                    if state and state.code == "33"
                    else Decimal("0"),
                    igst=Decimal("0") if state and state.code == "33" else round(gst, 2),
                    total_amount=round(total, 2),
                    remarks="Imported from Amazon order report",
                )

                for row, product in products_by_row:
                    quantity = as_int(row.data.get("quantity"))
                    gross_price = row_total(row)
                    unit_price = gross_price / Decimal(quantity) if quantity else Decimal("0")
                    OrderItem.objects.create(
                        order=order,
                        product=product,
                        quantity=quantity,
                        price=round(unit_price, 2),
                    )
                    result.item_count += 1
        except ValidationError as exc:
            result.skipped_count += 1
            result.skipped_orders.append(f"{order_number}: validation failed")
            result.errors.append(f"Order {order_number}: {exc}")
            continue

        result.created_count += 1
        result.created_orders.append(order_number)

    return result
