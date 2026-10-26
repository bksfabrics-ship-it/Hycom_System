from io import BytesIO

import pandas as pd
from django.test import SimpleTestCase
from openpyxl import load_workbook

from .invoice_converter import convert_dc_excel


class InvoiceConverterTests(SimpleTestCase):
    def _make_input(self):
        rows = [
            {
                "DC Number": "DC-101",
                "DC Date": "2026-10-01",
                "GST (if B2B Customer)": "Yes",
                "Portal": "Customer One",
                "SAP SKU": "SKU-001",
                "Qty": 2,
                "Taxable Amount": 100,
                "Tax": 5,
                "Shipping": None,
                "Total Amount": 105,
            },
            {
                "DC Number": "DC-102",
                "DC Date": "2026-10-02",
                "GST (if B2B Customer)": "No",
                "Portal": "Customer Two",
                "SAP SKU": "SKU-002",
                "Qty": 1,
                "Taxable Amount": 50,
                "Tax": 2.5,
                "Shipping": 10,
                "Total Amount": 62.5,
            },
        ]
        input_file = BytesIO()
        pd.DataFrame(rows).to_excel(input_file, index=False)
        input_file.seek(0)
        return input_file

    def test_converts_each_dc_to_a_formatted_worksheet(self):
        output = convert_dc_excel(self._make_input())
        workbook = load_workbook(output, data_only=True)

        self.assertEqual(workbook.sheetnames, ["DC_DC-101", "DC_DC-102"])
        first_sheet = workbook["DC_DC-101"]
        self.assertEqual(first_sheet["B4"].value, "B TO B")
        self.assertEqual(first_sheet["B7"].value, "DC-101")
        self.assertEqual(first_sheet["B9"].value, 2)
        self.assertEqual(first_sheet["C10"].value, 5)
        self.assertIsNone(first_sheet["E10"].value)
        self.assertEqual(first_sheet["B11"].value, "SKU-001")

        second_sheet = workbook["DC_DC-102"]
        self.assertEqual(second_sheet["B4"].value, "B TO C")
        self.assertEqual(second_sheet["E10"].value, 10)

    def test_accepts_missing_columns_and_preserves_rows_without_a_sku(self):
        input_file = BytesIO()
        pd.DataFrame(
            {"Qty": [3, "not numeric"], "Tax": [7, None]}
        ).to_excel(input_file, index=False)
        input_file.seek(0)

        workbook = load_workbook(convert_dc_excel(input_file), data_only=True)
        sheet = workbook[workbook.sheetnames[0]]

        self.assertIsNone(sheet["B3"].value)
        self.assertIsNone(sheet["B4"].value)
        self.assertIsNone(sheet["B5"].value)
        self.assertIsNone(sheet["B7"].value)
        self.assertEqual(sheet["B9"].value, 3)
        self.assertIsNone(sheet["B10"].value)
        self.assertEqual(sheet["D10"].value, 7)
        self.assertIsNone(sheet["E10"].value)
        self.assertIsNone(sheet["B11"].value)
        self.assertEqual(sheet["B13"].value, "not numeric")

    def test_keeps_rows_with_blank_dc_numbers_on_the_previous_dc_sheet(self):
        input_file = BytesIO()
        pd.DataFrame(
            {
                "DC Number": ["DC-101", None, None],
                "SAP SKU": ["SKU-001", "SKU-002", "SKU-003"],
                "Qty": [1, 2, 3],
            }
        ).to_excel(input_file, index=False)
        input_file.seek(0)

        workbook = load_workbook(convert_dc_excel(input_file), data_only=True)

        self.assertEqual(workbook.sheetnames, ["DC_DC-101"])
        sheet = workbook["DC_DC-101"]
        self.assertEqual(sheet["B7"].value, "DC-101")
        self.assertEqual(
            [sheet.cell(row=row, column=2).value for row in (9, 13, 17)],
            [1, 2, 3],
        )
        self.assertEqual(
            [sheet.cell(row=row, column=2).value for row in (11, 15, 19)],
            ["SKU-001", "SKU-002", "SKU-003"],
        )

    def test_converts_empty_workbook_to_a_blank_invoice_sheet(self):
        input_file = BytesIO()
        pd.DataFrame().to_excel(input_file, index=False)
        input_file.seek(0)

        workbook = load_workbook(convert_dc_excel(input_file), data_only=True)

        self.assertEqual(workbook.sheetnames, ["DC_"])
        self.assertIsNone(workbook["DC_"]["B3"].value)
